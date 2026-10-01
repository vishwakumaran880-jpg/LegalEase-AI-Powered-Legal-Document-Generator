from io import BytesIO
import re

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
)
from reportlab.lib import colors
from reportlab.lib.utils import ImageReader

def extract_terms(data: dict) -> list[dict]:
    fields = [
        ("Document Type", data.get("document_type", "")),
        ("Party A", data.get("party_a", "")),
        ("Party B", data.get("party_b", "")),
        ("Effective Date", data.get("effective_date", "")),
    ]
    if data.get("role"):
        fields.append(("Role", data["role"]))
    if data.get("compensation"):
        fields.append(("Compensation", data["compensation"]))
    if data.get("address"):
        fields.append(("Property Address", data["address"]))
    if data.get("lease_term"):
        fields.append(("Lease Term", data["lease_term"]))
    if data.get("monthly_rent"):
        fields.append(("Monthly Rent", data["monthly_rent"]))
    if data.get("confidentiality_scope"):
        fields.append(("Confidentiality Scope", data["confidentiality_scope"]))
    if data.get("key_terms"):
        fields.append(("Key Terms", data["key_terms"]))
    return [{"term": k, "value": v} for k, v in fields if v]

def _split_content(content: str):
    lines = [x.strip() for x in content.splitlines()]
    return [x for x in lines if x]

def build_docx(content, terms, company_name="", logo_bytes=None):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    if logo_bytes:
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(BytesIO(logo_bytes), width=Inches(1.2))
        except Exception:
            pass

    if company_name:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(company_name)
        r.bold = True
        r.font.size = Pt(14)

    lines = _split_content(content)
    for line in lines:
        p = doc.add_paragraph()
        if re.match(r"^(\d+[\.\)]|[A-Z][A-Z\s&-]{4,})", line):
            r = p.add_run(line)
            r.bold = True
        else:
            p.add_run(line)

    doc.add_page_break()
    doc.add_heading("Key Terms", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Term"
    table.rows[0].cells[1].text = "Value"
    for item in terms:
        cells = table.add_row().cells
        cells[0].text = str(item.get("term", ""))
        cells[1].text = str(item.get("value", ""))

    doc.add_paragraph(
        "AI-assisted draft — professional legal review recommended before signing."
    )

    out = BytesIO()
    doc.save(out)
    return out.getvalue()

def build_pdf(content, terms, company_name="", logo_bytes=None):
    out = BytesIO()
    doc = SimpleDocTemplate(
        out, pagesize=A4, rightMargin=45, leftMargin=45, topMargin=45, bottomMargin=45
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "LegalTitle", parent=styles["Title"], alignment=TA_CENTER, fontSize=15, leading=19
    )
    heading = ParagraphStyle(
        "LegalHeading", parent=styles["Heading2"], fontSize=11, leading=14, spaceBefore=7
    )
    body = ParagraphStyle(
        "LegalBody", parent=styles["BodyText"], fontSize=9.5, leading=13, spaceAfter=5
    )

    story = []
    if logo_bytes:
        try:
            img = RLImage(BytesIO(logo_bytes), width=70, height=70, kind="proportional")
            story += [img, Spacer(1, 6)]
        except Exception:
            pass

    if company_name:
        story.append(Paragraph(company_name, title))
        story.append(Spacer(1, 8))

    for line in _split_content(content):
        if re.match(r"^(\d+[\.\)]|[A-Z][A-Z\s&-]{4,})", line):
            story.append(Paragraph(line.replace("&", "&amp;"), heading))
        else:
            safe = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe, body))

    story.append(Spacer(1, 10))
    story.append(Paragraph("Key Terms", heading))
    table_data = [["Term", "Value"]] + [
        [str(x.get("term", "")), str(x.get("value", ""))] for x in terms
    ]
    table = Table(table_data, colWidths=[120, 340], repeatRows=1)
    table.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), 0.5, colors.grey),
        ("BACKGROUND", (0,0), (-1,0), colors.lightgrey),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("FONTSIZE", (0,0), (-1,-1), 8),
        ("LEFTPADDING", (0,0), (-1,-1), 5),
        ("RIGHTPADDING", (0,0), (-1,-1), 5),
    ]))
    story += [table, Spacer(1, 10)]
    story.append(Paragraph(
        "AI-assisted draft — professional legal review recommended before signing.",
        body
    ))

    doc.build(story)
    return out.getvalue()

def build_txt(content, terms, company_name=""):
    text = []
    if company_name:
        text.append(company_name)
        text.append("")
    text.append(content)
    text.append("")
    text.append("KEY TERMS")
    text.append("=" * 50)
    for item in terms:
        text.append(f'{item.get("term")}: {item.get("value")}')
    text.append("")
    text.append("AI-assisted draft — professional legal review recommended before signing.")
    return "\n".join(text).encode("utf-8")
